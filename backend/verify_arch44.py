"""ARCH-44 — Complex Table & Hierarchical Grid Extractor: verification harness.

Run from backend/:

    python verify_arch44.py                      # offline gates (engine, goldens, validator, exports, wiring, console)
    python verify_arch44.py --db                 # + extraction, schema refusals, HTTP 402/200, corrections, learned
                                                 #   mappings, hub, job, dispatch, erasure -- ONE rolled-back transaction
    python verify_arch44.py --mutate             # + deliberate breakages every gate must catch
    python verify_arch44.py --build              # + tsc -b, vite build, eslint on every ARCH-44 console file
    python verify_arch44.py --regression         # + verify_arch43.py --db
    python verify_arch44.py --db --mutate --build --regression   # certification
    python verify_arch44.py --write-goldens      # regenerate tests/fixtures/arch44_golden.json (review the diff!)

ARCH44-S1:verify. Evidence goes to backend/evidence/arch44/.
"""

from __future__ import annotations

import argparse
import contextlib
import hashlib
import importlib.util
import io
import json
import os
import random
import re
import subprocess
import sys
import time
import traceback
import uuid
import zipfile
from datetime import date
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Callable, Optional

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

BACKEND = Path(__file__).resolve().parent
ROOT = BACKEND.parent
FRONTEND = ROOT / "frontend"
SRC = FRONTEND / "src"
APP = BACKEND / "app"
VERSIONS = BACKEND / "alembic" / "versions"
EVIDENCE = BACKEND / "evidence" / "arch44"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

A43 = "arch43_step1_case_intelligence"
A44 = "arch44_step1_table_intelligence"
STEP3 = "arch40_step3_contract_ai_settings"
KEY = "capability.table_intelligence"
TABLES = ("extracted_tables", "extracted_table_cells", "table_validations", "table_column_mappings")
HELD_OUT_SEEDS = tuple(range(101, 111))


def read(path: Any) -> str:
    return Path(path).read_bytes().decode("utf-8-sig").replace("\r\n", "\n")


S = APP / "services" / "tables"
F = {
    "migration": VERSIONS / f"{A44}.py", "m43": VERSIONS / f"{A43}.py", "step3": VERSIONS / f"{STEP3}.py",
    "vocab": S / "vocabulary.py", "values": S / "values.py", "geometry": S / "geometry.py", "grid": S / "grid.py",
    "engine": S / "engine.py", "validate": S / "validate.py", "roles": S / "roles.py", "export": S / "export.py",
    "memory": S / "memory.py", "reader": S / "reader.py", "service": S / "service.py", "gate": S / "gate.py",
    "synthetic": S / "synthetic.py", "init": S / "__init__.py",
    "models": APP / "models/tables.py", "schemas": APP / "schemas/tables.py", "api": APP / "api/v1/tables.py",
    "handler": APP / "workers/handlers/tables.py",
    "ent": APP / "core/entitlements.py", "capgate": APP / "api/capability_gate.py", "seed": BACKEND / "scripts/seed_quota_tiers.py",
    "events": APP / "core/automation_events.py", "triggers": APP / "services/automation/triggers.py",
    "review_model": APP / "models/review.py", "review_vocab": APP / "services/review/vocabulary.py",
    "resolution": APP / "services/review/resolution.py", "review_schema": APP / "schemas/review.py",
    "review_api": APP / "api/v1/review.py", "router": APP / "api/v1/router.py",
    "handlers": APP / "workers/handlers/__init__.py", "profiles": APP / "workers/profiles.py",
    "post": APP / "services/post_enrichment.py", "erasure": APP / "services/compliance/erasure_service.py",
    "models_init": APP / "models/__init__.py", "conformance": BACKEND / "scripts/automation_conformance.py",
    "v36": BACKEND / "verify_arch36.py", "vhm": BACKEND / "verify_hardening_master.py",
    "backfill": BACKEND / "scripts/backfill_tables.py",
    "golden": BACKEND / "tests/fixtures/arch44_golden.json",
    "fe_types": SRC / "types/tables.ts", "fe_api": SRC / "services/api/tables.ts", "fe_list": SRC / "pages/tables/Tables.tsx",
    "fe_viewer": SRC / "pages/tables/TableViewer.tsx", "fe_doc": SRC / "components/tables/DocumentTables.tsx",
    "fe_caps": SRC / "constants/capabilities.ts", "fe_plan": SRC / "constants/planFeatures.ts",
    "fe_nav": SRC / "components/layout/navigation.ts", "fe_paths": SRC / "routes/tenantPaths.ts", "fe_app": SRC / "App.tsx",
    "fe_wid": SRC / "pages/WorkItems/WorkItemDetails.tsx", "fe_review_types": SRC / "types/review.ts",
    "fe_resolve": SRC / "components/review/ResolvePanel.tsx", "fe_hub": SRC / "pages/Verification/ReviewHub.tsx",
}
CHANGED_FRONTEND = tuple(str(F[k].relative_to(FRONTEND)).replace("\\", "/") for k in F if k.startswith("fe_"))


def t(key: str) -> str:
    return read(F[key])


class Recorder:
    def __init__(self) -> None:
        self.results: list[tuple[str, str, str]] = []

    def check(self, layer: str, name: str, fn: Callable[[], None]) -> bool:
        if ONLY and not any(name.startswith(prefix) for prefix in ONLY):
            return True
        try:
            fn()
        except AssertionError as exc:
            self.results.append((layer, name, f"FAIL  {exc}"))
            print(f"  FAIL  [{layer}] {name}\n        {str(exc)[:900]}")
            return False
        except Exception as exc:  # noqa: BLE001
            detail = "".join(traceback.format_exception_only(type(exc), exc)).strip()
            self.results.append((layer, name, f"ERROR {detail}"))
            print(f"  FAIL  [{layer}] {name}\n        {detail[:900]}")
            return False
        self.results.append((layer, name, "PASS"))
        print(f"  PASS  [{layer}] {name}")
        return True

    def summary(self) -> int:
        failed = [r for r in self.results if r[2] != "PASS"]
        by_layer: dict[str, int] = {}
        for layer, _, outcome in self.results:
            if outcome == "PASS":
                by_layer[layer] = by_layer.get(layer, 0) + 1
        print(f"\n{len(self.results) - len(failed)} passed, {len(failed)} failed ("
              + ", ".join(f"{n} {layer}" for layer, n in by_layer.items()) + ")")
        return 1 if failed else 0


#: --only G4,MS7,... runs just those gates (a sandbox with a per-call time
#: limit certifies in slices; on a workstation run everything in one command).
ONLY: tuple[str, ...] = ()


class AnchorMissing(RuntimeError):
    """A mutation whose anchor drifted. Never counted as 'caught'."""


def swap(text: str, old: str, new: str) -> str:
    if old not in text:
        raise AnchorMissing(f"mutation anchor missing: {old[:80]!r}")
    return text.replace(old, new, 1)


def expect_failure(fn: Callable[[], Any]) -> None:
    try:
        fn()
    except AnchorMissing:
        raise
    except Exception:  # noqa: BLE001
        return
    raise AssertionError("the gate PASSED against broken code")


def mutation(build: Callable[[], tuple], gate: Callable[..., Any]) -> Callable[[], None]:
    """Build the broken input FIRST, outside the refusal check."""
    def run() -> None:
        args = build()
        expect_failure(lambda: gate(*args))
    return run


def _load_module(name: str, path: Path, text: Optional[str] = None):
    spec = importlib.util.spec_from_file_location(name, path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module  # dataclasses and typing resolve the defining module through sys.modules
    if text is None:
        spec.loader.exec_module(module)
    else:
        exec(compile(text, str(path), "exec"), module.__dict__)  # noqa: S102
    return module


@contextlib.contextmanager
def patched(*items: tuple[Any, str, Any]):
    """Temporarily replace attributes; an attribute that does not exist is a drifted anchor."""
    saved = []
    for target, attr, value in items:
        if not hasattr(target, attr):
            raise AnchorMissing(f"{getattr(target, '__name__', target)}.{attr} missing")
        saved.append((target, attr, getattr(target, attr)))
        setattr(target, attr, value)
    try:
        yield
    finally:
        for target, attr, value in reversed(saved):
            setattr(target, attr, value)


def variant(key: str, old: str, new: str, attr: str):
    """One function of a module, rebuilt from its source with one change."""
    text = swap(t(key), old, new)
    module = _load_module(f"_mut_{key}_{abs(hash((old, new))) % 10**8}", F[key], text)
    return getattr(module, attr)


def _revisions() -> dict[str, str]:
    revs: dict[str, str] = {}
    for path in VERSIONS.glob("*.py"):
        text = read(path)
        r = re.search(r'^revision\s*(?::\s*str)?\s*=\s*["\']([^"\']+)', text, re.M)
        d = re.search(r'^down_revision\s*(?::[^=]+)?=\s*(.+)$', text, re.M)
        if r:
            revs[r.group(1)] = d.group(1).strip() if d else ""
    return revs


# ===========================================================================
# Engine helpers
# ===========================================================================


def run_case(case, *, options=None, raster: bool = True):
    from app.services.tables import engine as E
    from app.services.tables import reader

    pages = reader.pages_for(pdf_bytes=case.pdf, metadata={"pages": case.pages}, raster_rules=raster)
    return E.extract(pages, options=options)


def score(case, tables) -> dict:
    from app.services.tables import vocabulary as v

    truth = case.truth["tables"][0]
    out = {"case": case.name, "tables": len(tables)}
    if not tables:
        return {**out, "accuracy": 0.0, "exact": False}
    table = tables[0]
    grid = table.body_grid()
    cells = sum(len(r) for r in truth["rows"])
    hits = sum(1 for i, row in enumerate(truth["rows"]) for j, c in enumerate(row)
               if i < len(grid) and j < len(grid[i]) and grid[i][j] == c)
    flagged = sorted([0, c.row, c.col] for c in table.cells if v.FLAG_ARITH_FAIL in c.flags)
    out.update({
        "accuracy": round(hits / cells, 4), "cells": cells, "rows": [len(grid), len(truth["rows"])],
        "cols": [table.n_cols, len(truth["paths"])], "paths": [c.path for c in table.columns] == truth["paths"],
        "header_rows": table.header_rows == truth["header_rows"],
        "kinds": [r.kind for r in table.rows[table.header_rows:]] == truth["kinds"],
        "levels": [r.level for r in table.rows[table.header_rows:]] == truth["levels"],
        "pages": [table.page_start, table.page_end] == truth["pages"], "method": table.method,
        "method_ok": table.method == truth["method"], "flagged": flagged, "flags_ok": flagged == case.truth["flagged"],
        "status": table.status, "confidence": table.confidence, "rotation": table.rotation, "skew": table.skew,
    })
    out["exact"] = (len(tables) == 1 and out["accuracy"] == 1.0 and out["paths"] and out["header_rows"] and out["kinds"]
                    and out["levels"] and out["pages"] and out["method_ok"] and out["flags_ok"])
    return out


def assert_exact(result: dict) -> None:
    assert result.get("exact"), {k: v for k, v in result.items() if k not in ("flagged",) or not result.get("flags_ok")}


# ===========================================================================
# Offline gates
# ===========================================================================


def check_capability(ent: str, gate_text: str, seed: str, caps: str, plan: str, nav: str, v36: str, vhm: str) -> None:
    assert f'TABLE_INTELLIGENCE_CAPABILITY: str = "{KEY}"' in ent, "not declared in entitlements.py"
    block = ent.split("CAPABILITY_KEYS: tuple[str, ...] = (", 1)[1].split(")", 1)[0]
    assert "TABLE_INTELLIGENCE_CAPABILITY" in block, "not in CAPABILITY_KEYS"
    assert "name=TABLE_INTELLIGENCE_CAPABILITY" in ent, "no Entitlement entry (has_capability would raise)"
    assert "entitlements.TABLE_INTELLIGENCE_CAPABILITY:" in gate_text, "no 402 display name"
    business = seed.split("BUSINESS_CAPABILITIES = [", 1)[1].split("]", 1)[0]
    enterprise = seed.split("ENTERPRISE_CAPABILITIES = [", 1)[1].split("]", 1)[0]
    developer = seed.split("DEVELOPER_FEATURES = [", 1)[1].split("]", 1)[0]
    assert KEY in business and KEY in enterprise, "not packaged into Business and Enterprise"
    assert KEY not in developer, "leaked into Developer"
    assert f'tableIntelligence: "{KEY}"' in caps and "[CAPABILITY.tableIntelligence]:" in plan, "console capability or plan label missing"
    assert nav.count("capability: CAPABILITY.tableIntelligence") == 1, "nav entry not capability-locked"
    assert '"workspaceTables": "src/pages/tables/Tables.tsx"' in v36, "verify_arch36 GATED_PAGES not widened"
    assert 'BUS = BUS + ["capability.table_intelligence"]' in vhm and 'ENT = ENT + ["capability.table_intelligence"]' in vhm
    from app.core import entitlements

    assert KEY in entitlements.CAPABILITY_KEYS and KEY in entitlements.ENTITLEMENT_KEYS


def check_migration(text: str, revs: Optional[dict] = None) -> None:
    revs = revs if revs is not None else _revisions()
    assert revs.get(A44) == f'"{A43}"', f"{A44} revises {revs.get(A44)}"
    # ARCH45-S1:chain-widened-44. ARCH-45 sits between ARCH-44 and the contract step.
    assert revs.get(STEP3) in (f'"{A44}"', '"arch45_step1_corroboration"', '"arch46_step1_obligations"'), f"the contract step revises {revs.get(STEP3)}, expected {A44}, arch45 or arch46"  # ARCH46-S1:chain-widened-44
    if revs.get(STEP3) == '"arch45_step1_corroboration"':
        assert revs.get("arch45_step1_corroboration") == f'"{A44}"', "arch45 must revise arch44"
    if revs.get(STEP3) == '"arch46_step1_obligations"':  # ARCH46-S1:chain-widened-44
        assert revs.get("arch46_step1_obligations") == '"arch45_step1_corroboration"', "arch46 must revise arch45"
        assert revs.get("arch45_step1_corroboration") == f'"{A44}"', "arch45 must revise arch44"
    downs = " ".join(revs.values())
    heads = [r for r in revs if f'"{r}"' not in downs and f"'{r}'" not in downs]
    assert heads == [STEP3], f"file heads {heads}; the held contract step must stay the only head"
    for table in TABLES:
        assert f"CREATE TABLE {table}" in text, f"{table} missing"
    for name in ("fk_extracted_tables_work_item", "REFERENCES work_items (id, workspace_id) ON DELETE CASCADE",
                 "ck_extracted_tables_pages", "ck_extracted_tables_header", "ck_extracted_tables_confidence",
                 "ck_extracted_tables_flagged", "ck_extracted_tables_reviewed", "ck_extracted_tables_rows_len",
                 "ck_extracted_tables_cols_len", "pk_extracted_table_cells", "ck_extracted_table_cells_row",
                 "ck_extracted_table_cells_col", "ck_extracted_table_cells_span", "ck_extracted_table_cells_confidence",
                 "ck_extracted_table_cells_number", "ck_extracted_table_cells_date", "ck_extracted_table_cells_empty",
                 "ck_extracted_table_cells_flags", "CREATE TRIGGER trg_extracted_table_cells_within_grid",
                 "ck_extracted_table_cells_within_grid", "ck_extracted_table_cells_page_range",
                 "ck_table_validations_cell", "ck_table_validations_relation", "uq_table_column_mappings",
                 "fk_table_column_mappings_template", "'TABLE'::varchar(16)", "'TABLE_ARITHMETIC'::varchar(24)",
                 "ck_review_assignments_kind_known"):
        assert name in text, f"{name} missing from the migration"
    module = _load_module("_m44_check", F["migration"], text)
    m43 = _load_module("_m43_check", F["m43"])
    view = module.review_queue_view_v5()
    assert m43.review_queue_view_v4().rstrip() in view, "ARCH-44 altered an earlier arm of the hub view"
    assert module.REVIEW_KINDS[:5] == m43.REVIEW_KINDS and module.REVIEW_KINDS[5] == "TABLE"
    from app.core import automation_events as ae
    from app.models.review import REVIEW_KINDS
    from app.services.review import vocabulary as vocab
    from app.services.tables import vocabulary as tv

    # ARCH45-S1:vocab-widened-44. The newest hub migration defines the vocabulary;
    # ARCH-44's kinds and reasons must remain its prefix.
    newest45 = VERSIONS / "arch45_step1_corroboration.py"
    ref = _load_module("_m45_check44", newest45) if newest45.exists() else module
    newest46 = VERSIONS / "arch46_step1_obligations.py"  # ARCH46-S1:vocab-widened-44
    ref = _load_module("_m46_check44", newest46) if newest46.exists() else ref
    assert tuple(REVIEW_KINDS) == ref.REVIEW_KINDS and tuple(vocab.REASONS) == ref.REVIEW_REASONS, "hub vocabulary != migration"
    assert ref.REVIEW_KINDS[:len(module.REVIEW_KINDS)] == module.REVIEW_KINDS and ref.REVIEW_REASONS[:len(module.REVIEW_REASONS)] == module.REVIEW_REASONS
    for name in ("STATUSES", "METHODS", "VALUE_TYPES", "FLAGS", "CHECK_KINDS", "ROLES"):
        assert tuple(getattr(module, name)) == tuple(getattr(tv, name)), f"{name}: migration != vocabulary"
    assert set(module.NEW_TRIGGER_EVENTS) == set(ae.ARCH44_TRIGGER_EVENT_TYPES) <= set(ae.INTERNAL_EVENT_TYPES)
    internal = set(module.internal_after_44())
    if hasattr(ref, "internal_after_45"):  # ARCH45-S1:internal-widened-44 (adds trigger.corroboration.discrepancies)
        assert internal <= set(ref.internal_after_45())
        internal = set(ref.internal_after_45())
    if hasattr(ref, "internal_after_46"):  # ARCH46-S1:internal-widened-44 (adds the two obligation triggers)
        assert internal <= set(ref.internal_after_46())
        internal = set(ref.internal_after_46())
    assert set(ae.TRIGGER_NATIVE_EVENT_TYPES) | set(ae.TRIGGER_TWIN_EVENT_TYPES) <= internal, "a trigger event outside the outbox CHECK"


def check_dbscan() -> dict:
    """Our sweeps equal scikit-learn's DBSCAN(min_samples=1) on random inputs."""
    import numpy as np
    from sklearn.cluster import DBSCAN

    from app.services.tables import grid as G

    def same(a, b) -> bool:
        pairs = {}
        for x, y in zip(a, b):
            if pairs.setdefault(x, y) != y:
                return False
        return len(set(pairs.values())) == len(pairs)

    rng = random.Random(4401)
    for trial in range(120):
        n = rng.randint(1, 50)
        values = [rng.uniform(0, 200) for _ in range(n)]
        eps = rng.uniform(0.3, 9)
        theirs = DBSCAN(eps=eps, min_samples=1).fit(np.array(values).reshape(-1, 1)).labels_
        assert same(G.dbscan_1d(values, eps), list(theirs)), f"1-D trial {trial} differs"
        intervals = [(x, x + rng.uniform(0.5, 25)) for x in (rng.uniform(0, 300) for _ in range(n))]
        d = np.array([[max(0.0, max(a[0], b[0]) - min(a[1], b[1])) for b in intervals] for a in intervals])
        theirs = DBSCAN(eps=eps, min_samples=1, metric="precomputed").fit(d).labels_
        assert same(G.dbscan_intervals(intervals, eps), list(theirs)), f"interval trial {trial} differs"
    return {"trials": 120, "kinds": ["1-D row centres", "interval-gap metric"]}


def check_geometry() -> None:
    import ctypes
    import math

    import pikepdf
    import pypdfium2 as pdfium
    import pypdfium2.raw as raw

    from app.services.tables import geometry as g

    pdf = pikepdf.new()
    for rot in (0, 90, 180, 270):
        page = pdf.add_blank_page(page_size=(612, 792))
        if rot:
            page.Rotate = rot
    buf = io.BytesIO()
    pdf.save(buf)
    doc = pdfium.PdfDocument(buf.getvalue())
    rng = random.Random(4402)
    for i, rot in enumerate((0, 90, 180, 270)):
        page = doc[i]
        w, h = 612.0, 792.0
        dw, dh = (h, w) if rot in (90, 270) else (w, h)
        mapper = g.display_mapper(rot, w, h)
        for _ in range(20):
            x, y = rng.uniform(0, w), rng.uniform(0, h)
            dx, dy = ctypes.c_int(), ctypes.c_int()
            raw.FPDF_PageToDevice(page.raw, 0, 0, int(dw), int(dh), 0, ctypes.c_double(x), ctypes.c_double(y),
                                  ctypes.byref(dx), ctypes.byref(dy))
            mx, my = mapper(x, y)
            assert abs(mx - dx.value) <= 1 and abs(my - dy.value) <= 1, (rot, (x, y), (mx, my), (dx.value, dy.value))
    doc.close()
    # A quarter turn makes every reading direction read +x.
    for q, (vx, vy) in enumerate(((1, 0), (0, 1), (-1, 0), (0, -1))):
        assert g.quarter_of(vx, vy) == q
        frame = g.Frame(1000, 800, q)
        a, b = frame.turn(100, 100), frame.turn(100 + 50 * vx, 100 + 50 * vy)
        assert b[0] - a[0] > 0 and abs(b[1] - a[1]) < 1e-9, (q, a, b)
    # Deskew: points on a line tilted by theta come back level.
    theta = math.radians(1.7)
    frame = g.Frame(1700, 2200, 0, theta)
    ys = [frame.point(300 + k * 100, 500 + k * 100 * math.tan(theta))[1] for k in range(10)]
    assert max(ys) - min(ys) < 1e-6, ys
    assert abs(g.residual_skew([theta] * 5, 0) - theta) < 1e-12


VALUE_CASES = [
    ("1,23,456.78", "MONEY", "123456.78"), ("12,345.67", "MONEY", "12345.67"), ("1.234,56", "MONEY", "1234.56"),
    ("(1,250.00)", "MONEY", "-1250.00"), ("1,250.00-", "MONEY", "-1250.00"), ("-45.10", "MONEY", "-45.10"),
    ("5,000.00 Dr", "MONEY", "-5000.00"), ("5,000.00 Cr", "MONEY", "5000.00"), ("Rs. 1,000", "MONEY", "1000"),
    ("INR 2,500.50", "MONEY", "2500.50"), ("₹ 99.00", "MONEY", "99.00"), ("$1,234.00", "MONEY", "1234.00"),
    ("12.5%", "PERCENT", "12.5"), ("42", "NUMBER", "42"), ("3.5", "NUMBER", "3.5"), ("3.75", "MONEY", "3.75"),
    ("000012345678", "TEXT", None), ("50100123456789", "TEXT", None), ("13.0 - 17.0", "TEXT", None),
    ("-", "EMPTY", None), ("NIL", "EMPTY", None), ("12,34", "MONEY", "12.34"), ("1,000", "MONEY", "1000"),
]


def check_values() -> None:
    from app.services.tables import values as vals

    for text, kind, number in VALUE_CASES:
        parsed = vals.classify(text)
        assert parsed.kind == kind, (text, parsed.kind, kind)
        if number is not None:
            assert parsed.number == Decimal(number), (text, parsed.number, number)
    assert vals.parse_date("03/04/2026") == date(2026, 4, 3)
    assert vals.parse_date("03/04/2026", vals.ORDER_MDY) == date(2026, 3, 4)
    assert vals.parse_date("2026-05-03") == date(2026, 5, 3) and vals.parse_date("03-Apr-26") == date(2026, 4, 3)
    assert vals.parse_date("Apr 3, 2026") == date(2026, 4, 3) and vals.parse_date("31/02/2026") is None
    assert vals.infer_date_order(["04/05/2026", "04/28/2026"]) == vals.ORDER_MDY
    assert vals.infer_date_order(["04/05/2026", "28/04/2026"]) == vals.ORDER_DMY
    assert vals.infer_column_type(["04/05/2026", "04/28/2026", "05/13/2026"]) == ("DATE", vals.ORDER_MDY)
    assert vals.infer_column_type(["1,200.00", "35.50", "", "7.00"])[0] == "MONEY"
    parsed, ok = vals.parse_cell("abc", "MONEY")
    assert not ok and parsed.kind == "TEXT"


def golden_cases() -> list:
    from app.services.tables import synthetic as syn

    return syn.all_cases()


def golden_payload(cases) -> dict:
    from app.services.tables import vocabulary as v

    return {"_sentinel": "ARCH44-S1:golden", "engine_version": v.ENGINE_VERSION,
            "cases": {c.name: {"truth": c.truth, "pdf_sha256": hashlib.sha256(c.pdf).hexdigest(),
                               "source": c.source} for c in cases}}


def check_goldens(evidence: Optional[dict] = None) -> None:
    committed = json.loads(read(F["golden"]))
    cases = golden_cases()
    assert set(committed["cases"]) == {c.name for c in cases}, "golden case set changed"
    results, pdf_drift = [], []
    for case in cases:
        assert committed["cases"][case.name]["truth"] == json.loads(json.dumps(case.truth)), f"{case.name}: the generator's truth drifted from the golden file"
        if committed["cases"][case.name]["pdf_sha256"] != hashlib.sha256(case.pdf).hexdigest():
            pdf_drift.append(case.name)
        results.append(score(case, run_case(case)))
    if evidence is not None:
        evidence["golden"] = results
        evidence["golden_pdf_bytes_drift"] = pdf_drift
    bad = [r for r in results if not r["exact"]]
    assert not bad, bad[:3]


def check_held_out() -> dict:
    from app.services.tables import synthetic as syn

    makers = {"statement": lambda s, p: syn.statement(s, planted=p), "scan": lambda s, p: syn.scan(s, planted=p),
              "twolevel": lambda s, p: syn.twolevel(s, planted=p), "ledger": lambda s, p: syn.ledger(s, planted=p),
              "clinical": lambda s, p: syn.clinical(s), "lattice": lambda s, p: syn.lattice(s, planted=p),
              "lattice_scan": lambda s, p: syn.lattice_scan(s), "rotated90": lambda s, p: syn.rotated90(s),
              "sideways": lambda s, p: syn.sideways(s)}
    planted_ok = {"statement", "scan", "twolevel", "ledger", "lattice"}
    total, cells, bad = 0, 0, []
    for name, make in makers.items():
        for seed in HELD_OUT_SEEDS:
            for planted in ((False, True) if name in planted_ok else (False,)):
                case = make(seed, planted)
                result = score(case, run_case(case))
                total += 1
                cells += result.get("cells", 0)
                if not result["exact"]:
                    bad.append((name, seed, planted, {k: result.get(k) for k in ("accuracy", "paths", "kinds", "flags_ok", "tables")}))
    assert not bad, bad[:4]
    return {"documents": total, "cells": cells, "seeds": list(HELD_OUT_SEEDS), "exact": total}


def check_lattice() -> dict:
    from app.services.tables import engine as E
    from app.services.tables import synthetic as syn

    case = syn.lattice_scan()
    full = score(case, run_case(case))
    stream_only = score(case, run_case(case, options=E.Options(lattice=False)))
    assert full["exact"] and full["method"] == "LATTICE", full
    assert stream_only["accuracy"] < 0.8, f"stream alone should not separate merged OCR cells: {stream_only['accuracy']}"
    digital = score(syn.lattice(), run_case(syn.lattice()))
    assert digital["exact"] and digital["method"] == "LATTICE", digital
    return {"lattice_scan_accuracy": full["accuracy"], "stream_only_accuracy": stream_only["accuracy"]}


def check_continuation() -> dict:
    from app.services.tables import engine as E
    from app.services.tables import synthetic as syn

    case = syn.statement()
    tables = run_case(case)
    assert len(tables) == 1 and (tables[0].page_start, tables[0].page_end) == (1, 3) and tables[0].header_rows == 2
    assert tables[0].n_rows - tables[0].header_rows == len(case.truth["tables"][0]["rows"]), "repeated headers were kept"
    split = run_case(case, options=E.Options(continuation=False))
    assert len(split) == 3, f"without continuation the statement is {len(split)} tables"
    scan = run_case(syn.scan())
    assert len(scan) == 1 and (scan[0].page_start, scan[0].page_end) == (1, 2)
    return {"statement_pages": [1, 3], "without_continuation_tables": len(split)}


def check_rotation() -> dict:
    from app.services.tables import synthetic as syn

    side = score(syn.sideways(), run_case(syn.sideways()))
    turned = score(syn.rotated90(), run_case(syn.rotated90()))
    assert side["exact"] and side["rotation"] == 270, side
    assert turned["exact"], turned
    return {"sideways_rotation": side["rotation"], "rotated90_exact": turned["exact"]}


def check_deskew() -> dict:
    from app.services.tables import engine as E
    from app.services.tables import synthetic as syn

    case = syn.scan()
    fixed = score(case, run_case(case))
    raw = score(case, run_case(case, options=E.Options(deskew=False)))
    assert fixed["exact"] and abs(fixed["skew"] - 1.2) <= 0.2, fixed
    assert raw["accuracy"] < 0.5, f"a 1.2 degree scan should not survive without deskew: {raw['accuracy']}"
    return {"skew_corrected": fixed["skew"], "accuracy_without_deskew": raw["accuracy"]}


def check_validator() -> dict:
    from app.services.tables import synthetic as syn
    from app.services.tables import vocabulary as v

    out = {}
    for case in golden_cases():
        tables = run_case(case)
        flagged = sorted([0, c.row, c.col] for c in tables[0].cells if v.FLAG_ARITH_FAIL in c.flags)
        assert flagged == case.truth["flagged"], (case.name, flagged, case.truth["flagged"])
        out[case.name] = len(flagged)
    t = run_case(syn.statement(planted=True))[0]
    causes = {(c.row, c.col): c.detail.get("cause") for c in t.checks if c.scope == v.SCOPE_CELL}
    truth = syn.statement(planted=True).truth["flagged"]
    assert causes[(truth[0][1], truth[0][2])] == "AMOUNT" and causes[(truth[1][1], truth[1][2])] == "BALANCE", causes
    scan_t = run_case(syn.scan(planted=True))[0]
    carry = [c for c in scan_t.checks if c.kind == v.CHECK_CARRY_FORWARD and c.scope == v.SCOPE_CELL]
    assert len(carry) == 1, carry
    kinds = {name: sorted({c.kind for c in run_case(make())[0].checks if c.scope == v.SCOPE_RELATION})
             for name, make in (("statement", syn.statement), ("ledger", syn.ledger), ("lattice", syn.lattice),
                                ("clinical", syn.clinical), ("scan", syn.scan))}
    assert kinds["statement"] == ["COLUMN_SUM", "RUNNING_BALANCE"], kinds
    assert kinds["ledger"] == ["COLUMN_SUM", "HIERARCHY_SUM", "ROW_PRODUCT"], kinds
    assert kinds["lattice"] == ["COLUMN_SUM", "ROW_PRODUCT"] and kinds["clinical"] == [], kinds
    assert kinds["scan"] == ["RUNNING_BALANCE"], kinds
    return {"flagged_per_case": out, "relations": kinds}


def check_confidence() -> None:
    from app.services.tables import synthetic as syn
    from app.services.tables import validate as V
    from app.services.tables import vocabulary as v

    scan = run_case(syn.scan())[0]
    bal = next(c.index for c in scan.columns if c.role == v.ROLE_BALANCE)
    raised = [c for c in scan.cells if c.row >= scan.header_rows and c.col == bal and c.number is not None]
    assert raised and all(c.confidence > c.base_confidence for c in raised if c.base_confidence < 1), "passing checks did not raise confidence"
    planted = run_case(syn.statement(planted=True))[0]
    for c in planted.cells:
        if v.FLAG_ARITH_FAIL in c.flags:
            assert c.confidence == round(c.base_confidence * 0.4, 4), (c.row, c.col, c.confidence, c.base_confidence)
    assert V.adjust(0.8, 0, False) == 0.8 and V.adjust(0.8, 1, False) == 0.9 and V.adjust(0.8, 2, False) == 0.95
    assert V.adjust(0.8, 5, True) == 0.32


def check_exports() -> dict:
    from app.services.tables import engine as E
    from app.services.tables import export as X
    from app.services.tables import synthetic as syn

    planted = run_case(syn.statement(planted=True))[0]
    two = run_case(syn.twolevel())[0]
    text = X.to_csv(planted).decode("utf-8")
    assert text.startswith("﻿") and "\r\n" in text
    lines = text.lstrip("﻿").split("\r\n")
    assert lines[0] == "Date,Narration,Chq./Ref.No.,Value Dt,Withdrawal Amt.,Deposit Amt.,Closing Balance", lines[0]
    assert lines[1] == "2026-04-01,OPENING BALANCE,,,,,250000.00", lines[1]
    assert X.to_csv(two).decode("utf-8-sig").split("\r\n")[0] == "Date,Particulars,Amount / Debit,Amount / Credit,Balance"
    evil = E.OutTable(0, 1, 1, "STREAM", 0, 0.0, None, 2, 2, 1,
                      [E.OutColumn(0, ["Note"], "note", "note", "TEXT", "DMY", "OTHER", "INFERRED"),
                       E.OutColumn(1, ["Amount"], "amount", "amount", "MONEY", "DMY", "AMOUNT", "INFERRED")],
                      [E.OutRow(0, "HEADER", 0, 1), E.OutRow(1, "BODY", 0, 1)],
                      [E.OutCell(1, 0, 1, 1, 1, '=HYPERLINK("http://x")', "TEXT", None, None, 1, 1, 1, [], None),
                       E.OutCell(1, 1, 1, 1, 1, "-12.50", "MONEY", Decimal("-12.50"), None, 1, 1, 1, [], None)],
                      [], 1.0, "EXTRACTED", "0" * 40, [])
    row = X.to_csv(evil).decode("utf-8-sig").split("\r\n")[1]
    assert row == '"\'=HYPERLINK(""http://x"")",-12.50', row
    data = X.to_xlsx([planted, two])
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = set(z.namelist())
        assert {"[Content_Types].xml", "xl/workbook.xml", "xl/styles.xml", "xl/worksheets/sheet1.xml",
                "xl/worksheets/sheet2.xml", "_rels/.rels", "xl/_rels/workbook.xml.rels"} <= names
        sheet2 = z.read("xl/worksheets/sheet2.xml").decode()
        sheet1 = z.read("xl/worksheets/sheet1.xml").decode()
    assert '<mergeCell ref="C1:D1"/>' in sheet2 and '<mergeCell ref="A1:A2"/>' in sheet2, "header spans not merged"
    assert '<c r="G3" s="3"><v>250000.00</v></c>' in sheet1, "numbers are not numbers"
    assert '<c r="A3" s="2"><v>46113</v></c>' in sheet1, "dates are not date serials"
    assert 'state="frozen"' in sheet1 and 's="6"' in sheet1, "no frozen header or no shaded failures"
    readable = None
    try:
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb.worksheets[1]
        readable = {"sheets": wb.sheetnames, "merged": sorted(str(r) for r in ws.merged_cells.ranges)}
        assert wb.worksheets[0]["G3"].value == 250000 and "C1:D1" in readable["merged"]
    except ImportError:
        readable = "openpyxl not installed; structure checked with zipfile"
    doc = json.loads(X.to_json(planted))
    assert doc["records"][0]["closing_balance"] == "250000.00" and doc["records"][0]["date"] == "2026-04-01"
    assert doc["table"]["status"] == "FLAGGED" and any(v["scope"] == "CELL" for v in doc["validations"])
    return {"xlsx_bytes": len(data), "openpyxl": readable}


def check_memory() -> None:
    from app.services.tables import memory, roles
    from app.services.tables import vocabulary as v

    assert memory.decide([("DEBIT", 2, 0)]) is None, "two confirmations applied a mapping"
    assert memory.decide([("DEBIT", 3, 0)]) == "DEBIT"
    assert memory.decide([("DEBIT", 3, 1), ("CREDIT", 1, 3)]) is None, "a contradicted mapping was applied"
    assert memory.decide([("DEBIT", 5, 1)]) is None, "5 of 6 is below the bound (Wilson 0.497)"
    assert memory.decide([("DEBIT", 6, 1), ("CREDIT", 1, 6)]) == "DEBIT"
    samples = {("Withdrawal Amt.",): v.ROLE_DEBIT, ("Deposit Amt.",): v.ROLE_CREDIT, ("Closing Balance",): v.ROLE_BALANCE,
               ("Value Dt",): v.ROLE_VALUE_DATE, ("Chq./Ref.No.",): v.ROLE_REFERENCE, ("Amount", "Debit"): v.ROLE_DEBIT,
               ("Reference Range",): v.ROLE_RANGE, ("Unit Price",): v.ROLE_UNIT_PRICE, ("Qty",): v.ROLE_QUANTITY,
               ("Particulars",): v.ROLE_DESCRIPTION}
    for path, role in samples.items():
        assert roles.infer_role(list(path), "TEXT") == role, (path, roles.infer_role(list(path), "TEXT"), role)
    assert roles.infer_roles([["Mystery"]], ["MONEY"], {"mystery": "BALANCE"}) == [("BALANCE", v.ROLE_SOURCE_LEARNED)]
    assert roles.layout_key([["A"], ["B"]]) == roles.layout_key([["a"], ["b."]]) != roles.layout_key([["B"], ["A"]])


def check_wiring(texts: dict[str, str]) -> None:
    from app.services.automation import triggers
    from app.workers import handlers, profiles

    assert 'ARCH44_JOB_TYPES: frozenset[str] = frozenset({"tables.extract_document"})' in texts["handlers"]
    assert '"tables.extract_document": _tables_extract_document' in texts["handlers"] and "| ARCH44_JOB_TYPES" in texts["handlers"]
    light = texts["profiles"].split("LIGHT = WorkerProfile(", 1)[1].split("OCR = WorkerProfile(", 1)[0]
    ocr = texts["profiles"].split("OCR = WorkerProfile(", 1)[1].split("ENRICH = WorkerProfile(", 1)[0]
    assert '"tables.extract_document"' in ocr and '"tables.extract_document"' not in light, "extraction must run on the OCR profile"
    assert "tables.extract_document" in profiles.OCR.job_types and "tables.extract_document" not in profiles.LIGHT.job_types
    assert "tables.extract_document" in handlers._HANDLERS
    assert "table_gate.capability_held(db, organization_id)" in texts["post"] and "job_type=table_vocab.JOB_EXTRACT" in texts["post"], "dispatch missing or ungated"
    assert "_table_service.erase_for_work_items(db, work_item_ids)" in texts["erasure"], "ARCH-20 erasure keeps tables"
    spec = triggers.TRIGGERS_BY_KEY["table.flagged"]
    assert spec.capability == KEY and spec.has_document and spec.event_types == ("trigger.table.flagged",)
    assert (len(triggers.TRIGGERS), len(triggers.CATALOG_EVENT_TYPES)) in ((18, 19), (19, 20), (21, 22))  # ARCH45-S1:catalog-widened-44  ARCH46-S1:catalog-widened-44
    assert "emit_trigger(" in texts["service"] and "event_type=v.EVENT_TABLE_FLAGGED" in texts["service"]
    assert re.search(r"EXPECTED_TRIGGERS = (18|19|21)\b", texts["conformance"]), "the live conformance matrix does not expect table.flagged"  # ARCH45-S1:conformance-widened-44  ARCH46-S1:conformance-widened-44
    assert "vocab.KIND_TABLE: _resolve_table" in texts["resolution"]
    assert "    if TABLE_INTELLIGENCE_CAPABILITY in granted:\n        kinds.append(vocab.KIND_TABLE)\n" in texts["review_api"], "the hub shows TABLE without the capability"
    assert "table_verdict=body.table_verdict" in texts["review_api"] and "table_verdict: Optional[str] = None" in texts["review_schema"]
    assert "api_router.include_router(tables.router)" in texts["router"]
    assert "from app.models.tables import (" in texts["models_init"]
    assert "gate.capability_held(db, organization_id)" in texts["backfill"] and "job_type=v.JOB_EXTRACT" in texts["backfill"], \
        "the backfill must queue only for organizations holding the capability"


_ROUTE = re.compile(r'@router\.(get|put|post|patch|delete)\("([^"]+)"[^\n]*\n(?:async )?def (\w+)\(([\s\S]*?)\) -> [^:]+:\n([\s\S]*?)(?=\n\n\n@router|\n\n\n__all__)')


def check_api(api: str) -> None:
    routes = _ROUTE.findall(api)
    assert len(routes) == 9, f"expected 9 routes, found {len(routes)}"
    for method, path, name, sig, body in routes:
        first = [line.strip() for line in body.strip().splitlines()[:2]]
        assert first[0] == "_ws(context, workspace_id)" and first[1].startswith("_gate(db, context,"), f"{name} is not gated first"
        role = "RequireContributor" if method in ("put", "post", "patch", "delete") else "RequireViewer"
        assert f"Depends({role})" in sig, f"{name} ({method} {path}) needs {role}"


def _fields_py(text: str, cls: str) -> set[str]:
    body = text.split(f"class {cls}(BaseModel):", 1)[1].split("\n\n\n", 1)[0]
    return set(re.findall(r"^    (\w+):", body, re.M))


def _fields_ts(text: str, iface: str) -> set[str]:
    body = text.split(f"export interface {iface} {{", 1)[1].split("}", 1)[0]
    return set(re.findall(r"readonly (\w+)\??:", body))


def check_console(texts: dict[str, str]) -> None:
    for cls in ("TableColumn", "TableRowInfo", "TableCell", "TableValidationRow", "TableSummary", "TableList",
                "LearnedMapping", "TableDetail", "DocumentTables", "ExtractResult"):
        py, ts = _fields_py(texts["schemas"], cls), _fields_ts(texts["fe_types"], cls)
        assert py == ts, f"{cls}: API {sorted(py ^ ts)} differ"
    assert "capability: CAPABILITY.tableIntelligence" in texts["fe_nav"] and "tablesPath(orgSlug, workspaceSlug)" in texts["fe_nav"]
    assert '<Route path={ROUTE_PATTERNS.workspaceTables} element={<Tables />} />' in texts["fe_app"]
    assert '<Route path={ROUTE_PATTERNS.workspaceTable} element={<TableViewer />} />' in texts["fe_app"]
    assert 'workspaceTables: "tables"' in texts["fe_paths"] and 'workspaceTable: "tables/:tableId"' in texts["fe_paths"]
    viewer = texts["fe_viewer"]
    for marker in ("confidenceTone(cell.confidence)", "ring-2 ring-inset ring-red-500", "downloadTable(workspaceId, tableId, format",
                   "setColumnRole(", "correctCells(", "reviewTable(", "onDoubleClick={start}", "rowSpan={cell?.row_span ?? 1}",
                   "colSpan={cell?.col_span ?? 1}", "useCapabilityAccess(workspace?.organizationId ?? \"\", CAPABILITY.tableIntelligence)"):
        assert marker in viewer, f"viewer lacks {marker}"
    assert 'responseType: "blob"' in texts["fe_api"] and "apiClient.get<Blob>(`${one(workspaceId, tableId)}/export`" in texts["fe_api"], \
        "exports must go through the authenticated client"
    assert "<DocumentTables workItemId={workItem.id} />" in texts["fe_wid"]
    assert 'if (item.kind === "TABLE")' in texts["fe_resolve"] and 'onResolve({ table_verdict: "ACCEPT" })' in texts["fe_resolve"]
    assert '{ id: "TABLE", label: "Tables", kind: "TABLE" }' in texts["fe_hub"]
    assert '| "TABLE_ARITHMETIC"' in texts["fe_review_types"] and (
        '"SPLIT", "TABLE"]' in texts["fe_review_types"] or '"SPLIT", "TABLE", "CORROBORATION"]' in texts["fe_review_types"]
        or '"SPLIT", "TABLE", "CORROBORATION", "OBLIGATION"]' in texts["fe_review_types"])  # ARCH45-S1:console-kinds-widened-44  ARCH46-S1:console-kinds-widened-44


EDITED_OR_NEW = [k for k in F if k not in ("m43", "golden")]


def check_sentinels() -> None:
    missing = [str(F[k].relative_to(ROOT)) for k in EDITED_OR_NEW if "ARCH44-S" not in t(k)]
    assert not missing, f"no ARCH44 sentinel in {missing}"
    assert '"_sentinel": "ARCH44-S1:golden"' in t("golden")


def texts_all() -> dict[str, str]:
    return {k: t(k) for k in F if k not in ("golden",)}


def offline(rec: Recorder, evidence: dict) -> None:
    print("\nOffline")
    texts = texts_all()
    rec.check("offline", "T1 capability: entitlements (+Entitlement), 402 name, Business+Enterprise not Developer, console, nav lock, verify36, hardening matrix",
              lambda: check_capability(texts["ent"], texts["capgate"], texts["seed"], texts["fe_caps"], texts["fe_plan"],
                                       texts["fe_nav"], texts["v36"], texts["vhm"]))
    rec.check("offline", "T2 migration: 4 tables, grid trigger, CHECKs, arch43 -> arch44 -> contract, one head, hub view v5, vocabulary parity",
              lambda: check_migration(texts["migration"]))
    rec.check("offline", "G1 DBSCAN: row and column sweeps equal scikit-learn DBSCAN(min_samples=1) on 240 random inputs",
              lambda: evidence.__setitem__("dbscan", check_dbscan()))
    rec.check("offline", "G2 geometry: display mapping == FPDF_PageToDevice at 0/90/180/270; quarter turns read +x; deskew levels a tilted line",
              check_geometry)
    rec.check("offline", "G3 values: Indian/Western/European money, Dr/Cr, negatives, currency, identifiers, DMY/MDY inference, percent",
              check_values)
    rec.check("offline", "G4 golden files: 14 synthetic documents extracted exactly (every cell, header path, kind, level, page range, flag)",
              lambda: check_goldens(evidence))
    rec.check("offline", f"G5 held-out seeds {HELD_OUT_SEEDS[0]}-{HELD_OUT_SEEDS[-1]}: every generator, clean and planted, exact",
              lambda: evidence.__setitem__("held_out", check_held_out()))
    rec.check("offline", "G6 lattice: a scanned ruled table (rules from pixels, OCR merged across cells) is exact; stream alone is not",
              lambda: evidence.__setitem__("lattice", check_lattice()))
    rec.check("offline", "G7 continuation: a 3-page statement is ONE table, repeated headers dropped; without continuation it is 3",
              lambda: evidence.__setitem__("continuation", check_continuation()))
    rec.check("offline", "G8 rotation: /Rotate 90 and a sideways-printed table are exact; the sideways one reports 270",
              lambda: evidence.__setitem__("rotation", check_rotation()))
    rec.check("offline", "G9 deskew: a 1.2-degree scan is exact and reports its skew; without deskew it falls apart",
              lambda: evidence.__setitem__("deskew", check_deskew()))
    rec.check("offline", "V1 validator: planted errors flagged exactly (none on clean), balance vs amount localised, carry-forward, echoes explained",
              lambda: evidence.__setitem__("validator", check_validator()))
    rec.check("offline", "V2 confidence: passing checks raise it, failures cut it to 40%", check_confidence)
    rec.check("offline", "X1 exports: CSV (BOM, paths, typed, injection-safe), XLSX (merged header, numbers, dates, frozen, shaded), JSON",
              lambda: evidence.__setitem__("exports", check_exports()))
    rec.check("offline", "M1 learned mappings: Wilson-bounded (3 unanimous apply, 2 do not, contradictions hold back); keyword roles; layout key",
              check_memory)
    rec.check("offline", "W1 wiring: job on OCR profile, dispatch gated, erasure, trigger table.flagged (18/19), conformance 18, hub TABLE gated, router",
              lambda: check_wiring(texts))
    rec.check("offline", "W2 API: 9 routes, each gated first; writes CONTRIBUTOR, reads VIEWER", lambda: check_api(texts["api"]))
    rec.check("offline", "W3 console: type parity (10 models), locked nav, routes, viewer heat/failures/edit/roles/export, document panel, hub TABLE",
              lambda: check_console(texts))
    rec.check("offline", "S1 every ARCH-44 file carries its sentinel", check_sentinels)


# ===========================================================================
# Live database layer (one rolled-back transaction)
# ===========================================================================

PDF = "application/pdf"


def live_e2e(patches: Optional[list] = None) -> list[tuple[str, bool, str]]:
    import tempfile

    import sqlalchemy as sa
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from sqlalchemy.orm import Session

    from app.api import capability_gate, deps
    from app.api.v1 import tables as tables_api
    from app.api.v1.router import api_router
    from app.core import storage as storage_module
    from app.core.exception_handlers import domain_exception_handler
    from app.core.exceptions import FlowPilotError
    from app.core.storage import StorageNamespace, tenant_key
    from app.core.storage.local import LocalStorageDriver
    from app.db import session as session_module
    from app.db.session import engine
    from app.models.job import Job
    from app.models.tables import ExtractedTable, ExtractedTableCell, TableColumnMapping
    from app.models.work_item import WorkItem
    from app.services import audit_service, outbox_service, post_enrichment
    from app.services.compliance import erasure_service
    from app.services.review import projection, resolution
    from app.services.tables import service
    from app.services.tables import synthetic as syn
    from app.services.tables import vocabulary as v
    from app.workers import handlers as job_handlers

    job_handlers.register_all()
    Seeder = _load_module("_v40", BACKEND / "verify_arch40.py").Seeder
    steps: list[tuple[str, bool, str]] = []

    def step(name: str, fn, savepoint: bool = True) -> None:
        try:
            if savepoint:
                with db.begin_nested():
                    fn()
            else:
                fn()
            steps.append((name, True, ""))
        except Exception as exc:  # noqa: BLE001
            steps.append((name, False, f"{type(exc).__name__}: {exc}"[:900]))

    conn = engine.connect()
    outer = conn.begin()
    db = Session(bind=conn, join_transaction_mode="create_savepoint", expire_on_commit=False)
    held = {"value": True}
    emitted: list[tuple[str, Any]] = []
    denials: list = []
    tmp = tempfile.TemporaryDirectory(prefix="arch44-storage-")
    originals = [(capability_gate, "has_capability", capability_gate.has_capability),
                 (capability_gate, "granted_capabilities", capability_gate.granted_capabilities),
                 (audit_service, "record_independently", audit_service.record_independently),
                 (outbox_service, "emit_trigger", outbox_service.emit_trigger),
                 (storage_module, "_driver", storage_module._driver),
                 (session_module, "SessionLocal", session_module.SessionLocal)]
    capability_gate.has_capability = lambda *_a, capability_key=None, **_k: held["value"] and capability_key == KEY
    capability_gate.granted_capabilities = lambda *_a, **_k: [KEY] if held["value"] else []
    audit_service.record_independently = lambda **kw: denials.append(kw)  # the org exists only in this transaction
    real_emit = outbox_service.emit_trigger

    def recording_emit(*a, **kw):
        emitted.append((kw.get("event_type"), kw.get("idempotency_key")))
        return real_emit(*a, **kw)

    outbox_service.emit_trigger = recording_emit
    storage_module._driver = LocalStorageDriver(root=Path(tmp.name))

    class _Scoped:
        def __call__(self):
            return self

        def __enter__(self):
            return db

        def __exit__(self, *exc):
            return False

    session_module.SessionLocal = _Scoped()
    for target, attr, value in patches or []:
        originals.append((target, attr, getattr(target, attr)))
        setattr(target, attr, value)
    s: dict[str, Any] = {}
    try:
        seed = Seeder(conn)
        org, ws, ws2, user = uuid.uuid4(), uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
        seed.insert("organizations", id=org, name="arch44", slug=f"arch44-{org.hex[:8]}", status="ACTIVE")
        seed.insert("users", id=user, email=f"t-{org.hex[:8]}@arch44.test", is_active=True, is_superuser=False,
                    is_verified=True, timezone="UTC", locale="en")
        for wid, name in ((ws, "tables-a"), (ws2, "tables-b")):
            seed.insert("workspaces", id=wid, organization_id=org, workspace_name=name, slug=f"{name}-{org.hex[:6]}",
                        status="ACTIVE", timezone="UTC", language="en", currency="USD", date_format="YYYY-MM-DD")
        seed.insert("workspace_members", id=uuid.uuid4(), user_id=user, workspace_id=ws, role="ADMIN", status="ACTIVE")

        def document(case, name: str, workspace=ws) -> WorkItem:
            wid = uuid.uuid4()
            key = tenant_key(organization_id=org, namespace=StorageNamespace.DOCUMENTS, file_id=wid, suffix="pdf")
            storage_module._driver.put(key, case.pdf, PDF)
            seed.insert("work_items", id=wid, workspace_id=workspace, original_filename=name, stored_filename=key,
                        file_type=PDF, file_size=len(case.pdf), page_count=len(case.pages),
                        extracted_text="\n".join(p.get("text", "") for p in case.pages),
                        extraction_metadata=json.dumps({"pages": case.pages}), extracted_entities=json.dumps({}),
                        created_by_user_id=user, pipeline_stage="COMPLETED")
            return db.get(WorkItem, wid)

        def refused(fn) -> bool:
            try:
                with db.begin_nested():
                    fn()
                    db.flush()
            except Exception:  # noqa: BLE001
                return True
            return False

        # ------------------------------------------------------------------ D2 drift
        def d2() -> None:
            from alembic.autogenerate import compare_metadata
            from alembic.migration import MigrationContext

            import app.models  # noqa: F401
            from app.db.base import Base

            diff = compare_metadata(MigrationContext.configure(conn, opts={"compare_type": True}), Base.metadata)
            flat = []
            for d in diff:
                flat.extend(d if isinstance(d, list) else [d])
            bad = []
            for d in flat:
                if d[0] not in ("modify_type", "add_column", "remove_column", "modify_nullable", "add_table", "remove_table"):
                    continue
                names = {getattr(x, "name", None) for x in d[1:] if hasattr(x, "name")} | {x for x in d[1:] if isinstance(x, str)}
                table_obj = next((x for x in d[1:] if x.__class__.__name__ == "Table"), None)
                if names & set(TABLES) or (table_obj is not None and table_obj.name in TABLES):
                    bad.append(str(d)[:160])
            assert not bad, bad

        step("D2 models match the migration: zero column-level drift on the four ARCH-44 tables", d2, savepoint=False)

        # ------------------------------------------------------------------ D3 refusals
        def d3() -> None:
            item = document(syn.clinical(), "refusals.pdf")
            table = ExtractedTable(id=uuid.uuid4(), organization_id=org, workspace_id=ws, work_item_id=item.id, ordinal=0,
                                   page_start=1, page_end=1, n_rows=3, n_cols=2, header_rows=1, method="STREAM",
                                   columns=[{}, {}], rows=[{}, {}, {}], bboxes=[], confidence=Decimal("0.9"),
                                   status="EXTRACTED", layout_key="a" * 40, engine_version=v.ENGINE_VERSION)
            db.add(table)
            db.flush()

            def cell(**kw):
                base = dict(table_id=table.id, workspace_id=ws, row_index=1, col_index=0, page=1, text="x",
                            value_type="TEXT", base_confidence=Decimal("0.9"), confidence=Decimal("0.9"), flags=[])
                base.update(kw)
                return ExtractedTableCell(**base)

            checks = {
                "cell outside the grid (trigger)": lambda: db.add(cell(row_index=3)),
                "span past the grid (trigger)": lambda: db.add(cell(col_index=1, col_span=2)),
                "page outside the table (trigger)": lambda: db.add(cell(page=2)),
                "negative row": lambda: db.add(cell(row_index=-1)),
                "confidence above 1": lambda: db.add(cell(confidence=Decimal("1.5"))),
                "a number on a TEXT cell": lambda: db.add(cell(value_number=Decimal("1"))),
                "EMPTY with text": lambda: db.add(cell(value_type="EMPTY")),
                "unknown flag": lambda: db.add(cell(flags=["BOGUS"])),
                "cross-workspace cell": lambda: db.add(cell(workspace_id=ws2)),
                "FLAGGED with no failures": lambda: db.execute(sa.update(ExtractedTable).where(ExtractedTable.id == table.id).values(status="FLAGGED")),
                "REVIEWED without a reviewer time": lambda: db.execute(sa.update(ExtractedTable).where(ExtractedTable.id == table.id).values(status="REVIEWED")),
                "rows array length != n_rows": lambda: db.execute(sa.update(ExtractedTable).where(ExtractedTable.id == table.id).values(n_rows=5)),
                "bad layout key": lambda: db.execute(sa.update(ExtractedTable).where(ExtractedTable.id == table.id).values(layout_key="z" * 40)),
                "relation outcome inconsistent": lambda: db.execute(sa.text(
                    "INSERT INTO table_validations (id, table_id, workspace_id, kind, scope, outcome, checked, failed) "
                    "VALUES (:i, :t, :w, 'ROW_TOTAL', 'RELATION', 'PASS', 3, 1)"), {"i": uuid.uuid4(), "t": table.id, "w": ws}),
                "cell failure without a cell": lambda: db.execute(sa.text(
                    "INSERT INTO table_validations (id, table_id, workspace_id, kind, scope, outcome) "
                    "VALUES (:i, :t, :w, 'ROW_TOTAL', 'CELL', 'FAIL')"), {"i": uuid.uuid4(), "t": table.id, "w": ws}),
                "mapping with an unknown role": lambda: db.add(TableColumnMapping(id=uuid.uuid4(), organization_id=org, workspace_id=ws,
                                                                                  layout_key="a" * 40, header_key="h", role="NOPE")),
            }
            accepted = [name for name, fn in checks.items() if not refused(fn)]
            assert not accepted, f"the schema accepted: {accepted}"
            assert not refused(lambda: db.add(cell(row_index=2, col_index=1, text="5", value_type="NUMBER", value_number=Decimal(5)))), \
                "a valid cell was refused"

        step("D3 schema refusals: grid/page trigger, spans, CHECKs on confidence/type/value/flags, FK, status rules, validations, roles", d3)

        # ------------------------------------------------------------------ D4 extraction
        def d4() -> None:
            clean = document(syn.statement(), "statement.pdf")
            result = service.extract_document(db, work_item=clean)
            tables = service.tables_of(db, clean.id)
            assert result["tables"] == 1 and len(tables) == 1 and tables[0].status == "VALIDATED", (result, [t.status for t in tables])
            out = service.load(db, tables[0])
            assert out.body_grid() == syn.statement().truth["tables"][0]["rows"], "the stored grid differs from the truth"
            assert (tables[0].page_start, tables[0].page_end, tables[0].header_rows) == (1, 3, 2)
            before = len([e for e in emitted if e[0] == v.EVENT_TABLE_FLAGGED])
            assert before == 0, "a table that reconciles emitted table.flagged"
            case = syn.statement(planted=True)
            planted = document(case, "statement-planted.pdf")
            service.extract_document(db, work_item=planted)
            table = service.tables_of(db, planted.id)[0]
            fails = sorted([0, r, c] for r, c in db.execute(sa.text(
                "SELECT row_index, col_index FROM table_validations WHERE table_id = :t AND scope = 'CELL'"), {"t": table.id}))
            assert table.status == "FLAGGED" and table.failed_checks == 3 and fails == case.truth["flagged"], (table.status, fails)
            flagged_events = [e for e in emitted if e[0] == v.EVENT_TABLE_FLAGGED]
            assert len(flagged_events) == 1, f"table.flagged emitted {len(flagged_events)} times"
            service.set_column_role(db, table=table, col=1, role="DESCRIPTION", actor_user_id=user)
            assert len([e for e in emitted if e[0] == v.EVENT_TABLE_FLAGGED]) == 1, "re-validation re-emitted table.flagged"
            hub = db.execute(sa.text("SELECT status, review_reason FROM review_queue_items WHERE kind = 'TABLE' AND item_id = :i"),
                             {"i": table.id}).all()
            assert hub == [("OPEN", "TABLE_ARITHMETIC")], hub
            scan_case = syn.scan()
            scanned = document(scan_case, "scan.pdf")
            service.extract_document(db, work_item=scanned)
            st = service.tables_of(db, scanned.id)[0]
            assert service.load(db, st).body_grid() == scan_case.truth["tables"][0]["rows"] and float(st.skew_degrees) > 1.0
            s.update(clean=clean, planted=planted, table=table, case=case)

        step("D4 extraction through the service: stored grid == truth (digital and scanned), FLAGGED at the planted cells, "
             "table.flagged once (not on re-validation), hub item OPEN", d4, savepoint=False)

        # ------------------------------------------------------------------ HTTP
        ctx = SimpleNamespace(workspace_id=ws, organization_id=org, user_id=user)
        app = FastAPI()
        app.include_router(api_router, prefix="/api/v1")
        app.add_exception_handler(FlowPilotError, domain_exception_handler)
        app.dependency_overrides[deps.get_db] = lambda: db
        for name in ("RequireViewer", "RequireContributor"):
            app.dependency_overrides[getattr(tables_api, name)] = lambda: ctx
        client = TestClient(app)
        s["client"] = client
        base = f"/api/v1/workspaces/{ws}"

        def d5() -> None:
            table, planted, case = s["table"], s["planted"], s["case"]
            held["value"] = False
            routes = (("get", "/tables"), ("get", f"/tables/{table.id}"), ("get", f"/tables/{table.id}/export?format=csv"),
                      ("patch", f"/tables/{table.id}/cells"), ("put", f"/tables/{table.id}/columns/0/role"),
                      ("post", f"/tables/{table.id}/review"), ("get", f"/work-items/{planted.id}/tables"),
                      ("post", f"/work-items/{planted.id}/tables/extract"), ("get", f"/work-items/{planted.id}/tables/export"))
            bodies = {"patch": {"cells": [{"row": 2, "col": 0, "text": "x"}]}, "put": {"role": "DATE"}, "post": {"verdict": "ACCEPT"}}
            for method, path in routes:
                kwargs = {"json": bodies[method]} if method in bodies else {}
                r = getattr(client, method)(base + path, **kwargs)
                assert r.status_code == 402 and "CAPABILITY" in r.text.upper(), (path, r.status_code, r.text[:200])
            assert any(d.get("details", {}).get("capability_key") == KEY for d in denials), "402 without the denial audit"
            held["value"] = True
            listed = client.get(base + "/tables").json()
            assert listed["counts_by_status"]["FLAGGED"] >= 1 and listed["total"] >= 3, listed["counts_by_status"]
            detail = client.get(base + f"/tables/{table.id}").json()
            assert len(detail["columns"]) == 7 and detail["table"]["header_rows"] == 2
            assert any(c["row_span"] == 2 for c in detail["cells"] if c["row"] == 0), "stacked header labels lost their row span"
            assert len([x for x in detail["validations"] if x["scope"] == "CELL"]) == 3
            for fmt, ctype in (("csv", "text/csv"), ("xlsx", "spreadsheetml"), ("json", "application/json")):
                r = client.get(base + f"/tables/{table.id}/export", params={"format": fmt})
                assert r.status_code == 200 and ctype in r.headers["content-type"] and len(r.content) > 100, (fmt, r.status_code)
            assert client.get(base + f"/tables/{table.id}/export", params={"format": "pdf"}).status_code == 422
            wb = client.get(base + f"/work-items/{planted.id}/tables/export")
            assert wb.status_code == 200 and wb.content[:2] == b"PK"
            exported = db.execute(sa.text("SELECT count(*) FROM audit_logs WHERE organization_id = :o AND action = 'EXPORTED'"),
                                  {"o": org}).scalar_one()
            assert exported >= 4, f"exports were not audited ({exported})"
            doc = client.get(base + f"/work-items/{planted.id}/tables").json()
            assert doc["extractable"] and len(doc["tables"]) == 1
            assert client.put(base + f"/tables/{table.id}/columns/0/role", json={"role": "NOPE"}).status_code == 422
            assert client.patch(base + f"/tables/{table.id}/cells", json={"cells": [{"row": 999, "col": 0, "text": "x"}]}).status_code == 422
            fixes = [{"row": r, "col": c, "text": syn.inr(Decimal(x["expected"]))} for (_, r, c), x in
                     zip(case.truth["flagged"], sorted([x for x in detail["validations"] if x["scope"] == "CELL"],
                                                       key=lambda x: (x["row_index"], x["col_index"])))]
            corrected = client.patch(base + f"/tables/{table.id}/cells", json={"cells": fixes})
            assert corrected.status_code == 200, corrected.text[:300]
            body = corrected.json()
            assert body["table"]["status"] == "VALIDATED" and body["table"]["failed_checks"] == 0, body["table"]
            fixed = [c for c in body["cells"] if "CORRECTED" in c["flags"]]
            assert len(fixed) == 3 and all(c["original_text"] and c["confidence"] == 1.0 for c in fixed)
            hub = db.execute(sa.text("SELECT status FROM review_queue_items WHERE kind = 'TABLE' AND item_id = :i"), {"i": table.id}).all()
            assert hub == [("RESOLVED",)], f"a table that now reconciles is still open in the hub: {hub}"
            again = client.post(base + f"/work-items/{planted.id}/tables/extract")
            assert again.status_code == 409 and "KEPT_CORRECTIONS" in again.text, (again.status_code, again.text[:200])
            forced = client.post(base + f"/work-items/{planted.id}/tables/extract", params={"force": "true"})
            assert forced.status_code == 200 and forced.json()["ran"] == "SYNC", forced.text[:200]
            new_id = forced.json()["tables"][0]["id"]
            assert client.post(base + f"/tables/{new_id}/review", json={"verdict": "ACCEPT"}).json()["table"]["status"] == "REVIEWED"
            assert client.post(base + f"/tables/{new_id}/review", json={"verdict": "REJECT"}).status_code == 409

        step("D5 HTTP: 402 on all 9 routes (+denial audit); list/detail/exports (audited)/document; 422s; corrections -> VALIDATED "
             "and hub RESOLVED; re-extract 409 then force; review 200 then 409", d5, savepoint=False)

        # ------------------------------------------------------------------ D6 learned mappings
        def d6() -> None:
            docs = [document(syn.twolevel(seed), f"ledger-{seed}.pdf") for seed in (201, 202, 203, 204)]
            for item in docs[:3]:
                service.extract_document(db, work_item=item)
            first = service.tables_of(db, docs[0].id)[0]
            assert first.columns[4]["role"] == "BALANCE" and first.columns[4]["role_source"] == "INFERRED"
            # Reviewers say "Balance" is a running TOTAL here (a deliberate relabel the keywords would never infer).
            for n, item in enumerate(docs[:3]):
                table = service.tables_of(db, item.id)[0]
                service.set_column_role(db, table=table, col=4, role="TOTAL", actor_user_id=user)
                if n == 1:
                    probe = document(syn.twolevel(299), "probe.pdf")
                    service.extract_document(db, work_item=probe)
                    assert service.tables_of(db, probe.id)[0].columns[4]["role"] == "BALANCE", "applied after two confirmations"
            service.extract_document(db, work_item=docs[3])
            learned = service.tables_of(db, docs[3].id)[0].columns[4]
            assert (learned["role"], learned["role_source"]) == ("TOTAL", "LEARNED"), learned
            row = db.execute(sa.select(TableColumnMapping).where(TableColumnMapping.workspace_id == ws,
                                                                 TableColumnMapping.role == "TOTAL")).scalar_one()
            assert row.confirmations == 3 and row.header_key == "balance"

        step("D6 learned column mapping: not applied after 2 confirmations, applied (LEARNED) to the next same-layout table after 3", d6)

        # ------------------------------------------------------------------ D7 hub
        def d7() -> None:
            case = syn.ledger(planted=True)
            item = document(case, "ledger-planted.pdf")
            service.extract_document(db, work_item=item)
            table = service.tables_of(db, item.id)[0]
            assert table.status == "FLAGGED"
            loaded = projection.load_item(db, workspace_id=ws, kind="TABLE", item_id=table.id)
            assert loaded is not None and loaded.status == "OPEN"
            outcome = resolution.resolve_item(db, item=loaded, actor_user_id=user,
                                              payload=resolution.ResolvePayload(table_verdict="ACCEPT"))
            db.refresh(table)
            assert table.status == "REVIEWED" and table.reviewed_by_user_id == user, (table.status, outcome)
            again = projection.load_item(db, workspace_id=ws, kind="TABLE", item_id=table.id)
            assert again is not None and again.status == "RESOLVED"
            from app.api.v1 import review as review_api

            held["value"] = False
            assert "TABLE" not in review_api._allowed_kinds(db, ctx)
            held["value"] = True
            assert "TABLE" in review_api._allowed_kinds(db, ctx)

        step("D7 review hub: a flagged table is a TABLE item; ACCEPT through resolution -> REVIEWED/RESOLVED; kind hidden without the plan", d7)

        # ------------------------------------------------------------------ D8 job and dispatch
        def d8() -> None:
            item = document(syn.lattice(), "invoice.pdf")
            before = db.execute(sa.select(sa.func.count()).select_from(Job).where(Job.job_type == v.JOB_EXTRACT)).scalar_one()
            post_enrichment.dispatch_in_session(db, work_item_id=item.id, organization_id=org, workspace_id=ws)
            after = db.execute(sa.select(sa.func.count()).select_from(Job).where(Job.job_type == v.JOB_EXTRACT)).scalar_one()
            assert after == before + 1, "enrichment did not enqueue tables.extract_document"
            held["value"] = False
            other = document(syn.clinical(), "chart.pdf")
            post_enrichment.dispatch_in_session(db, work_item_id=other.id, organization_id=org, workspace_id=ws)
            assert db.execute(sa.select(sa.func.count()).select_from(Job).where(Job.job_type == v.JOB_EXTRACT)).scalar_one() == after, \
                "dispatched without the capability"
            result = job_handlers._HANDLERS[v.JOB_EXTRACT]({"work_item_id": str(item.id)})
            assert result.get("extracted") is False and "not held" in result.get("reason", ""), result
            held["value"] = True
            result = job_handlers._HANDLERS[v.JOB_EXTRACT]({"work_item_id": str(item.id)})
            assert result.get("extracted") and result["tables"] == 1, result
            assert service.tables_of(db, item.id)[0].method == "LATTICE"

        step("D8 job: enrichment enqueues tables.extract_document only with the plan; the handler extracts (LATTICE) only with the plan",
             d8, savepoint=False)  # the handler commits its own session

        # ------------------------------------------------------------------ D9 erasure
        def d9() -> None:
            item = document(syn.ledger(), "erase-me.pdf")
            service.extract_document(db, work_item=item)
            assert service.tables_of(db, item.id)
            counts: dict[str, int] = {}
            erasure_service._destroy_documents(db, subject_id=user, workspace_ids=[ws], counts=counts)
            remaining = db.execute(sa.text("SELECT count(*) FROM extracted_tables WHERE workspace_id = :w"), {"w": ws}).scalar_one()
            cells = db.execute(sa.text("SELECT count(*) FROM extracted_table_cells WHERE workspace_id = :w"), {"w": ws}).scalar_one()
            assert remaining == 0 and cells == 0 and counts.get("extracted_tables", 0) >= 1, (remaining, cells, counts)

        step("D9 ARCH-20 erasure: a subject's documents lose every extracted table and cell", d9)
    finally:
        for target, attr, value in reversed(originals):
            setattr(target, attr, value)
        if "client" in s:
            s["client"].close()
        db.close()
        outer.rollback()
        conn.close()
        tmp.cleanup()
    return steps


def db_layer(rec: Recorder, evidence: dict, mutate: bool) -> None:
    print("\nDatabase")
    import sqlalchemy as sa

    def head() -> None:
        from app.db.session import engine

        with engine.connect() as conn:
            current = [r[0] for r in conn.execute(sa.text("SELECT version_num FROM alembic_version"))]
            tables = {r[0] for r in conn.execute(sa.text("SELECT table_name FROM information_schema.tables WHERE table_schema='public'"))}
            kinds = conn.execute(sa.text("SELECT pg_get_constraintdef(oid) FROM pg_constraint WHERE conname='ck_review_assignments_kind_known'")).scalar_one()
            outbox = conn.execute(sa.text("SELECT pg_get_constraintdef(oid) FROM pg_constraint WHERE conname='ck_outbox_events_visibility_vocabulary'")).scalar_one()
            triggers = {r[0] for r in conn.execute(sa.text("SELECT tgname FROM pg_trigger WHERE NOT tgisinternal"))}
        assert current in ([A44], [STEP3], ["arch45_step1_corroboration"], ["arch46_step1_obligations"]), f"alembic current is {current}; run run_arch44.ps1"  # ARCH45-S1:head-widened-44  ARCH46-S1:head-widened-44
        assert not [x for x in TABLES if x not in tables], "ARCH-44 tables missing"
        assert "TABLE" in kinds and "trigger.table.flagged" in outbox, (kinds[-80:], outbox[-120:])
        assert "trg_extracted_table_cells_within_grid" in triggers

    if not rec.check("db", "D1 database at arch44: 4 tables, the grid trigger, TABLE in the hub CHECK, table.flagged in the outbox CHECK", head):
        return
    base_run = not ONLY or any(p[0] == "D" for p in ONLY)
    results = live_e2e() if base_run else []
    if base_run:
        evidence["live"] = [{"step": n_, "ok": ok, "detail": d} for n_, ok, d in results]
    for name, ok, detail in results:
        rec.check("db", name, (lambda: None) if ok else (lambda d=detail: (_ for _ in ()).throw(AssertionError(d))))
    if not mutate:
        return

    from app.api.v1 import tables as api_module
    from app.services.compliance import erasure_service
    from app.services.review import resolution
    from app.services.tables import engine as engine_module
    from app.services.tables import memory as memory_module
    from app.services.tables import service as service_module

    def must_fail(step_prefix: str, build: Callable[[], list]) -> Callable[[], None]:
        def run() -> None:
            patches = build()  # anchors first: a drifted anchor raises AnchorMissing, never "caught"
            for target, attr, _ in patches:
                if not hasattr(target, attr):
                    raise AnchorMissing(f"{getattr(target, '__name__', target)}.{attr} missing")
            steps = live_e2e(patches)
            named = [ok for n_, ok, _ in steps if n_.startswith(step_prefix)]
            assert named, f"no step {step_prefix}"
            assert not all(named), f"{step_prefix} passed against broken code"
        return run

    def always_emit():
        return variant("service", "    if table.status == v.STATUS_FLAGGED and not was_flagged:\n        filename = db.execute(select(WorkItem.original_filename).where(WorkItem.id == table.work_item_id)).scalar_one()\n        _emit_flagged(db, table, filename)\n    _audit(db, table, \"column_role\"",
                       "    if table.status == v.STATUS_FLAGGED:\n        filename = db.execute(select(WorkItem.original_filename).where(WorkItem.id == table.work_item_id)).scalar_one()\n        _emit_flagged(db, table, filename)\n    _audit(db, table, \"column_role\"", "set_column_role")

    def no_revalidate():
        return variant("service", "    was_flagged = table.status == v.STATUS_FLAGGED\n    out = load(db, table)\n    E.revalidate(out)\n",
                       "    was_flagged = table.status == v.STATUS_FLAGGED\n    out = load(db, table)\n", "correct_cells")

    def overwrite():
        return variant("service", "    if existing and not force and any(t.corrected_at or t.status in v.DECIDED_STATUSES for t in existing):",
                       "    if False:", "extract_document")

    rec.check("mutation", "MD1 the table API's capability gate removed (D5 must catch)",
              must_fail("D5", lambda: [(api_module, "_gate", lambda *a, **k: None)]))
    rec.check("mutation", "MD2 table.flagged emitted on every re-validation, not once per entry (D4 must catch)",
              must_fail("D4", lambda: [(service_module, "set_column_role", always_emit())]))
    rec.check("mutation", "MD3 the hub cannot resolve TABLE (D7 must catch)", must_fail("D7", lambda: [
        (resolution, "_DISPATCH", {k: fn for k, fn in resolution._DISPATCH.items() if k != "TABLE"})]))
    rec.check("mutation", "MD4 a correction does not re-validate the table (D5 must catch)",
              must_fail("D5", lambda: [(service_module, "correct_cells", no_revalidate())]))
    rec.check("mutation", "MD5 learned mappings never applied (D6 must catch)",
              must_fail("D6", lambda: [(memory_module, "learned_for", lambda *a, **k: {})]))
    rec.check("mutation", "MD6 ARCH-20 erasure keeps the tables (D9 must catch)",
              must_fail("D9", lambda: [(service_module, "erase_for_work_items", lambda db, ids: 0)]))
    rec.check("mutation", "MD7 re-extraction overwrites a person's corrections without force (D5 must catch)",
              must_fail("D5", lambda: [(service_module, "extract_document", overwrite())]))
    rec.check("mutation", "MD8 dispatch ignores the plan (D8 must catch)", must_fail("D8", lambda: [
        (importlib.import_module("app.services.tables.gate"), "capability_held", lambda *a, **k: True)]))
    rec.check("mutation", "MD9 validation disabled in the engine (D4 must catch)", must_fail("D4", lambda: [
        (engine_module, "revalidate", lambda table, enabled=True: engine_module.V.Result())]))
    del erasure_service


# ===========================================================================
# Static mutations
# ===========================================================================


def mutations() -> list[tuple[str, Callable[[], None]]]:
    from app.services.tables import engine as E
    from app.services.tables import export as X
    from app.services.tables import geometry as geo
    from app.services.tables import grid as G
    from app.services.tables import memory as M
    from app.services.tables import reader
    from app.services.tables import synthetic as syn
    from app.services.tables import validate as V
    from app.services.tables import values as vals
    from app.services.tables import vocabulary as v

    texts = texts_all()

    def texts_with(key: str, text: str) -> dict:
        out = dict(texts)
        out[key] = text
        return out

    def cap(index: int, old: str, new: str) -> tuple:
        args = [texts["ent"], texts["capgate"], texts["seed"], texts["fe_caps"], texts["fe_plan"], texts["fe_nav"],
                texts["v36"], texts["vhm"]]
        args[index] = swap(args[index], old, new)
        return tuple(args)

    def under(patches: list, gate: Callable[[], Any]) -> Callable[[], None]:
        def run() -> None:
            for target, attr, _ in patches:
                if not hasattr(target, attr):
                    raise AnchorMissing(f"{getattr(target, '__name__', target)}.{attr} missing")
            with patched(*patches):
                expect_failure(gate)
        return run

    def statement_seed_5() -> None:
        case = syn.statement(5)
        assert_exact(score(case, run_case(case)))

    migration = texts["migration"]
    return [
        ("MS1 capability missing from the Business tier", mutation(lambda: cap(2, '    _capability("capability.table_intelligence"),  # ARCH44-S1:tier-business\n', ""), check_capability)),
        ("MS2 no Entitlement entry (has_capability would raise)", mutation(lambda: cap(0, "name=TABLE_INTELLIGENCE_CAPABILITY", "name=CASE_INTELLIGENCE_CAPABILITY"), check_capability)),
        ("MS3 console plan label missing (tsc would fail)", mutation(lambda: cap(4, "[CAPABILITY.tableIntelligence]:", "[CAPABILITY.caseIntelligence]:"), check_capability)),
        ("MS4 the grid trigger dropped", mutation(lambda: (swap(migration, "CREATE TRIGGER trg_extracted_table_cells_within_grid", "-- no trigger"),), check_migration)),
        ("MS5 the cell confidence CHECK dropped", mutation(lambda: (swap(migration, "CONSTRAINT ck_extracted_table_cells_confidence CHECK (", "CONSTRAINT ck_confidence_gone CHECK ("),), check_migration)),
        ("MS6 contract step left on arch43 (two heads)", mutation(lambda: (migration, {**_revisions(), STEP3: f'"{A43}"'}), check_migration)),
        ("MS7 column DBSCAN eps x8 (adjacent columns merge)", under([(v, "COL_EPS_H", v.COL_EPS_H * 8)], lambda: check_goldens())),
        ("MS8 figures side by side merged into one phrase", under([(G, "phrases_of", variant("grid", "        if out and ((gap <= gap_h * h and not figures) or", "        if out and ((gap <= gap_h * h) or", "phrases_of"))], statement_seed_5)),
        ("MS9 continuation across pages disabled", under([(E, "column_mapping", lambda a, b: None)], check_continuation)),
        ("MS10 reading-direction vote disabled (sideways tables)", under([(geo, "dominant_quarter", lambda d: 0)], check_rotation)),
        ("MS11 deskew disabled", under([(geo, "residual_skew", lambda angles, quarter: 0.0)], check_deskew)),
        ("MS12 lattice grids disabled", under([(G, "lattice_tables", lambda page, tokens, h: ([], set()))], check_lattice)),
        ("MS13 group labels centred over a column gap treated as single-column", under([(G, "header_place", G.place)], lambda: check_goldens())),
        ("MS14 TOTAL detection by prefix ('Total WBC Count' becomes a total)", under([(v, "TOTAL_RE", re.compile(r"^\s*(grand\s+|net\s+)?totals?\b", re.I))], lambda: check_goldens())),
        ("MS15 balance-misread localisation removed", under([(V, "_running_balance", variant("validate", "        if nxt is not None and (i + 1) in fails and nxt.prev_r == t.r and close(t.prev_b + t.delta + nxt.delta, nxt.b):", "        if False:", "_running_balance"))], check_validator)),
        ("MS16 echoes flagged: a total failing only because of a flagged cell", under([(V, "_sum_check", variant("validate", "        if close(corrected, printed):\n            continue  # explained by a cell already flagged\n", "", "_sum_check"))], check_validator)),
        ("MS17 tolerance loosened to 1,000", under([(V, "close", lambda a, b, tol=None: abs(a - b) <= Decimal("1000"))], check_validator)),
        ("MS18 confidence raise inverted", under([(V, "adjust", lambda base, support, flagged: round(base * (0.4 if flagged else 0.5), 4))], check_confidence)),
        ("MS19 Indian digit grouping not recognised", under([(vals, "_INT_INDIAN", re.compile(r"^$"))], check_values)),
        ("MS20 day/month order never inferred from the column", under([(vals, "infer_date_order", lambda texts_, default=vals.ORDER_DMY: vals.ORDER_DMY)], check_values)),
        ("MS21 CSV formula-injection guard removed", under([(X, "safe_text", lambda text: text)], check_exports)),
        ("MS22 XLSX header spans not merged", under([(X, "_sheet_xml", variant("export", "            if cell.row_span > 1 or cell.col_span > 1:\n", "            if False:\n", "_sheet_xml"))], check_exports)),
        ("MS23 Wilson bound bypassed (one confirmation applies a mapping)", under([(v, "MAPPING_MIN_SUPPORT", 1), (v, "MAPPING_APPLY_WILSON", 0.0)], check_memory)),
        ("MS24 raster rules at a fixed dark threshold (hairlines lost)", under([(reader, "rules_from_bitmap",
            lambda gray, px_to_frame=1.0, dark=None, min_len_px=None: geo.rules_from_bitmap(gray, px_to_frame=px_to_frame, dark=110))], check_lattice)),
        ("MS25 extraction moved to the LIGHT profile", mutation(lambda: (texts_with("profiles", swap(texts["profiles"], '            "tables.extract_document",\n', "").replace(
            '            "cases.assemble_document",\n', '            "cases.assemble_document",\n            "tables.extract_document",\n', 1)),), check_wiring)),
        ("MS26 post-enrichment dispatch ungated", mutation(lambda: (texts_with("post", swap(texts["post"], "and table_gate.capability_held(db, organization_id)", "and True")),), check_wiring)),
        ("MS27 ARCH-20 erasure hook removed", mutation(lambda: (texts_with("erasure", swap(texts["erasure"], "_table_service.erase_for_work_items(db, work_item_ids)", "0")),), check_wiring)),
        ("MS28 the hub shows TABLE without the plan", mutation(lambda: (texts_with("review_api", swap(texts["review_api"], "    if TABLE_INTELLIGENCE_CAPABILITY in granted:\n        kinds.append(vocab.KIND_TABLE)\n", "    kinds.append(vocab.KIND_TABLE)\n")),), check_wiring)),
        ("MS29 an API route ungated", mutation(lambda: (swap(texts["api"], '    _gate(db, context, "tables.export")\n', ""),), check_api)),
        ("MS30 a write route open to viewers", mutation(lambda: (swap(texts["api"], "                        context: TenantContext = Depends(RequireContributor)) -> TableDetail:\n    _ws(context, workspace_id)\n    _gate(db, context, \"tables.correct\")",
                                                                          "                        context: TenantContext = Depends(RequireViewer)) -> TableDetail:\n    _ws(context, workspace_id)\n    _gate(db, context, \"tables.correct\")"),), check_api)),
        ("MS31 console type drift (a TableSummary field removed)", mutation(lambda: (texts_with("fe_types", swap(texts["fe_types"], "  readonly failed_checks: number;\n  readonly checked_relations", "  readonly checked_relations")),), check_console)),
        ("MS32 exports fetched without the authenticated client", mutation(lambda: (texts_with("fe_api", swap(texts["fe_api"], 'responseType: "blob", timeout: 120_000,\n  });\n  downloadBlob(response.data, `${stem(filename)}_table', 'timeout: 120_000,\n  });\n  downloadBlob(response.data, `${stem(filename)}_table').replace('responseType: "blob"', "")),), check_console)),
        ("MS33 the resolve panel has no TABLE branch", mutation(lambda: (texts_with("fe_resolve", swap(texts["fe_resolve"], 'if (item.kind === "TABLE")', 'if (item.kind === ("NONE" as string))')),), check_console)),
        ("MS34 the trigger emitter removed", mutation(lambda: (texts_with("service", swap(texts["service"], "event_type=v.EVENT_TABLE_FLAGGED", "event_type=None")),), check_wiring)),
    ]


def _run(cmd: list[str], cwd: Path, timeout: int = 1800) -> None:
    proc = subprocess.run(cmd, cwd=str(cwd), capture_output=True, text=True, timeout=timeout,
                          shell=os.name == "nt", encoding="utf-8", errors="replace")
    if proc.returncode != 0:
        raise AssertionError(f"{' '.join(cmd)} exited {proc.returncode}\n{(proc.stdout + proc.stderr)[-2500:]}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Verify ARCH-44")
    for flag in ("--db", "--mutate", "--build", "--regression", "--write-goldens"):
        parser.add_argument(flag, action="store_true")
    parser.add_argument("--only", default="", help="comma-separated gate-name prefixes (e.g. G4,MS7)")
    parser.add_argument("--evidence", default="verify_arch44.json", help="evidence file name under evidence/arch44/")
    args = parser.parse_args()
    global ONLY
    ONLY = tuple(p.strip() for p in args.only.split(",") if p.strip())
    if args.write_goldens:
        payload = golden_payload(golden_cases())
        F["golden"].parent.mkdir(parents=True, exist_ok=True)
        F["golden"].write_text(json.dumps(payload, indent=1, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8")
        print(f"wrote {F['golden'].relative_to(ROOT)} ({len(payload['cases'])} cases)")
        return 0
    rec = Recorder()
    evidence: dict[str, Any] = {"started": time.strftime("%Y-%m-%dT%H:%M:%S"), "python": sys.version.split()[0]}
    t0 = time.perf_counter()
    offline(rec, evidence)
    if args.db:
        db_layer(rec, evidence, args.mutate)
    if args.mutate:
        print("\nMutations")
        for name, fn in mutations():
            rec.check("mutation", name, fn)
    if args.build:
        print("\nBuild")
        rec.check("build", "B1 tsc -b and vite build", lambda: (_run(["npx", "tsc", "-b"], FRONTEND), _run(["npx", "vite", "build"], FRONTEND)))
        rec.check("build", f"B2 eslint --max-warnings=0 on the {len(CHANGED_FRONTEND)} ARCH-44 console files",
                  lambda: _run(["npx", "eslint", "--max-warnings=0", *CHANGED_FRONTEND], FRONTEND))
    if args.regression:
        print("\nRegression")
        rec.check("regression", "R1 verify_arch43.py --db", lambda: _run([sys.executable, "verify_arch43.py", "--db"], BACKEND, 3600))
    evidence["seconds"] = round(time.perf_counter() - t0, 1)
    evidence["results"] = [{"layer": layer, "gate": name, "outcome": outcome} for layer, name, outcome in rec.results]
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    (EVIDENCE / args.evidence).write_text(json.dumps(evidence, indent=2, default=str), encoding="utf-8")
    return rec.summary()


if __name__ == "__main__":
    sys.exit(main())
